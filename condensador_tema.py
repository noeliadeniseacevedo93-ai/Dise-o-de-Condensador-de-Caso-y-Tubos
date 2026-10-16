#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AGENTE DE DISEÑO DE CONDENSADOR DE CASCO Y TUBOS (TEMA)
=======================================================
Fase I   : Módulo de entrada de datos (diccionario / archivo JSON)
Fase II  : Algoritmo iterativo térmico e hidráulico
Fase III : Exportación de la Hoja de Especificaciones TEMA (.xlsx)

Unidades: SI (m, kg, s, W, Pa, K/°C). Calor latente en J/kg, presiones en Pa.

Uso:
    python condensador_tema.py                      # corre el caso ejemplo
    python condensador_tema.py --ejemplo datos.json # guarda un JSON de ejemplo para editar
    python condensador_tema.py --entrada datos.json --salida hoja_TEMA.xlsx
    python condensador_tema.py --entrada datos.json --editar  # abre los datos en $EDITOR para modificarlos
"""
import argparse
import json
import math

G = 9.81  # m/s2

# ----------------------------------------------------------------------------
# TABLAS NORMALIZADAS
# ----------------------------------------------------------------------------
# Largos comerciales de tubo (m): 6, 8, 10, 12, 16, 20 ft
L_STD = [1.83, 2.44, 3.05, 3.66, 4.88, 6.10]

# Diámetros internos nominales de casco (pulgadas), serie TEMA / tubería
DS_STD = [x * 0.0254 for x in
          [8, 10, 12, 13.25, 15.25, 17.25, 19.25, 21.25, 23.25, 25, 27, 29, 31,
           33, 35, 37, 39, 42, 45, 48, 54, 60]]

# Conteo de tubos: Nt = K1 * (Db/Do)^n1   (Towler & Sinnott, Pt = 1.25 Do)
# clave = número de pasos por tubo
K_TRI = {1: (0.319, 2.142), 2: (0.249, 2.207), 4: (0.175, 2.285),
         6: (0.0743, 2.499), 8: (0.0365, 2.675)}
K_CUA = {1: (0.215, 2.207), 2: (0.156, 2.291), 4: (0.158, 2.263),
         6: (0.0402, 2.617), 8: (0.0331, 2.643)}

# Holgura casco - haz (m) según el tipo de cabezal posterior (última letra TEMA)
HOLGURA = {"M": 0.012, "N": 0.012, "L": 0.012, "U": 0.015,
           "S": 0.060, "T": 0.095, "P": 0.050, "W": 0.035}


# ----------------------------------------------------------------------------
# FASE I: ENTRADAS (ejemplo)
# ----------------------------------------------------------------------------
EJEMPLO = {
    "identificacion": {
        "item": "E-101",
        "servicio": "Condensador de vapor de agua",
        "cliente": "Práctica Integradora",
        "tema": "BEM",              # cabezal frontal / casco / cabezal posterior
        "orientacion": "Horizontal",
        "material_tubos": "Acero al carbono",
        "material_casco": "Acero al carbono",
    },
    "caliente": {                   # Lado que condensa
        "fluido": "Vapor de agua saturado",
        "Wh": 2.0,                  # kg/s
        "T1": 120.2, "T2": 120.2,   # °C (entrada / salida)
        "T_sat": 120.2,             # °C
        "P_oper": 200e3,            # Pa abs
        "lambda": 2201.9e3,         # J/kg
        "cp_vapor": 0.0, "cp_liquido": 0.0,   # J/kg K (solo si hay desrecalentamiento / subenfriamiento)
        "rho_liquido": 942.9, "mu_liquido": 2.32e-4, "k_liquido": 0.685,
        "rho_vapor": 1.129, "mu_vapor": 1.3e-5,
        "dp_adm": 20e3,             # Pa
        "Rd": 0.00009,              # m2 K/W
    },
    "fria": {
        "fluido": "Agua de enfriamiento",
        "Wc": None,                 # kg/s (None = se calcula del balance)
        "t1": 25.0, "t2": 40.0,
        "P_oper": 400e3,
        "cp": 4178.0, "mu": 7.6e-4, "k": 0.620, "rho": 995.0,   # a temperatura media
        "dp_adm": 70e3,
        "Rd": 0.00035,
    },
    "geometria": {
        "Do": 0.01905, "Di": 0.01483,   # m (3/4" BWG 14)
        "L_max": 4.88,                  # m
        "arreglo": "triangular",        # "triangular" (60°/30°) o "cuadro" (90°)
        "Pt": 0.02381,                  # m
        "k_pared": 50.0,                # W/m K
        "Ud_est": 2000.0,               # W/m2 K, estimación inicial
        "v_min": 1.0, "v_max": 3.0,     # m/s lado fluido sensible
    },
    "condensa_en": "casco",             # "casco" o "tubos"
}


# ----------------------------------------------------------------------------
# UTILIDADES TÉRMICAS
# ----------------------------------------------------------------------------
def dmlt_contracorriente(T1, T2, t1, t2):
    d1, d2 = T1 - t2, T2 - t1
    if d1 <= 0 or d2 <= 0:
        raise ValueError("Cruce de temperaturas: T1 <= t2 o T2 <= t1.")
    if abs(d1 - d2) < 1e-9:
        return d1
    return (d1 - d2) / math.log(d1 / d2)


def factor_ft(T1, T2, t1, t2, Ns):
    """Ft (Bowman) para Ns cascos en serie, 2+ pasos por tubo en cada casco."""
    R = (T1 - T2) / (t2 - t1)
    P = (t2 - t1) / (T1 - t1)
    if R < 1e-6:
        return 1.0            # un fluido isotérmico: Ft = 1
    if abs(R - 1) < 1e-6:
        R = 1 + 1e-6
    if P >= 1 or 1 - P * R <= 0:
        return 0.0
    F = ((1 - P * R) / (1 - P)) ** (1.0 / Ns)
    P1 = (1 - F) / (R - F)
    S = math.sqrt(R * R + 1)
    a = (1 - P1) / (1 - R * P1)
    b = (2 - P1 * (R + 1 - S)) / (2 - P1 * (R + 1 + S))
    if a <= 0 or b <= 0:
        return 0.0
    return S * math.log(a) / ((R - 1) * math.log(b))


def nusselt_tubo(Re, Pr):
    def gnielinski(Re):
        f = (0.79 * math.log(Re) - 1.64) ** -2
        return (f / 8) * (Re - 1000) * Pr / (1 + 12.7 * math.sqrt(f / 8) * (Pr ** (2 / 3) - 1))
    if Re >= 4000:
        return gnielinski(Re)
    if Re <= 2300:
        return 3.66
    w = (Re - 2300) / 1700.0
    return 3.66 * (1 - w) + gnielinski(4000) * w


def factor_friccion(Re, Di, rug=4.6e-5):
    if Re < 2300:
        return 64.0 / Re
    inv = -1.8 * math.log10((rug / Di / 3.7) ** 1.11 + 6.9 / Re)
    return 1.0 / inv ** 2


def kern_casco(W, rho, mu, k, cp, Ds, B, Pt, Do, L, arreglo):
    """Coeficiente y caída de presión lado casco (Kern, segmentado 25 %)."""
    C = Pt - Do
    As = Ds * C * B / Pt
    Gs = W / As
    if arreglo == "cuadro":
        De = 4 * (Pt ** 2 - math.pi * Do ** 2 / 4) / (math.pi * Do)
    else:
        De = 4 * (0.5 * Pt * 0.866 * Pt - 0.5 * math.pi * Do ** 2 / 4) / (math.pi * Do / 2)
    Re = De * Gs / mu
    Pr = cp * mu / k if cp else 1.0
    hs = 0.36 * k / De * Re ** 0.55 * Pr ** (1 / 3) if k else 0.0
    f = math.exp(0.576 - 0.19 * math.log(Re))
    Nb = max(1, int(round(L / B - 1)))
    dP = f * Gs ** 2 * Ds * (Nb + 1) / (2 * rho * De)
    return dict(hs=hs, dP=dP, Re=Re, Gs=Gs, De=De, Nb=Nb, As=As, v=Gs / rho)


def h_cond_horizontal(k, rho, rhov, mu, lam, dT, Do, Nr):
    """Nusselt para haz de tubos horizontales, corregido por filas (Kern)."""
    return 0.725 * (k ** 3 * rho * (rho - rhov) * G * lam / (mu * dT * Do)) ** 0.25 * Nr ** (-1 / 6)


def h_cond_vertical(k, rho, rhov, mu, lam, dT, L):
    """Nusselt película laminar para tubo vertical."""
    return 0.943 * (k ** 3 * rho * (rho - rhov) * G * lam / (mu * dT * L)) ** 0.25


def capacidad_tubos(Ds, n, Do, Pt, arreglo, rear):
    K1, n1 = (K_CUA if arreglo == "cuadro" else K_TRI)[n]
    Db = Ds - HOLGURA.get(rear, 0.012)
    if Db <= 0:
        return 0
    return int(K1 * (Db / Do) ** n1 * (1.25 * Do / Pt) ** 2)


# ----------------------------------------------------------------------------
# FASE II: PREPARACIÓN, EVALUACIÓN Y BÚSQUEDA ITERATIVA
# ----------------------------------------------------------------------------
def preparar(d, warns):
    h, c = d["caliente"], d["fria"]
    h["Tsat"] = h.get("T_sat") or (h["T1"] if abs(h["T1"] - h["T2"]) < 1e-9 else h["T1"])
    Q = h["Wh"] * (h["lambda"]
                   + h.get("cp_vapor", 0) * max(h["T1"] - h["Tsat"], 0)
                   + h.get("cp_liquido", 0) * max(h["Tsat"] - h["T2"], 0))
    Wc_bal = Q / (c["cp"] * (c["t2"] - c["t1"]))
    if c.get("Wc"):
        if abs(c["Wc"] - Wc_bal) / Wc_bal > 0.02:
            warns.append(f"Wc ingresado ({c['Wc']:.2f}) difiere del balance ({Wc_bal:.2f} kg/s). "
                         "Se usa el valor del balance de energía.")
    c["Wc"] = Wc_bal
    d["Q"] = Q
    geo = d["geometria"]
    d["Rd_req"] = (h["Rd"] + c["Rd"] * geo["Do"] / geo["Di"]) if d["condensa_en"] == "casco" \
        else (c["Rd"] + h["Rd"] * geo["Do"] / geo["Di"])
    d["rear"] = d["identificacion"]["tema"].strip().upper()[-1]


def evaluar(d, g):
    h, c, geo = d["caliente"], d["fria"], d["geometria"]
    Do, Di, Pt, kw = geo["Do"], geo["Di"], geo["Pt"], geo["k_pared"]
    arr = geo["arreglo"]
    Ns, Nt, n, L, Ds, B = g["Ns"], g["Nt"], g["n"], g["L"], g["Ds"], g["B"]
    Tsat = h["Tsat"]
    Tc = 0.5 * (c["t1"] + c["t2"])
    Rw = Do * math.log(Do / Di) / (2 * kw)
    A = Ns * Nt * math.pi * Do * L
    r = dict(A=A, Ns=Ns, Nt=Nt, n=n, L=L, Ds=Ds, B=B)

    if d["condensa_en"] == "casco":
        # ---- Tubos: agua (sensible)
        Ap = Nt / n * math.pi * Di ** 2 / 4
        Gt = c["Wc"] / Ap
        Re = Gt * Di / c["mu"]
        Pr = c["cp"] * c["mu"] / c["k"]
        hi = nusselt_tubo(Re, Pr) * c["k"] / Di
        f = factor_friccion(Re, Di)
        dPt = Ns * f * (L * n / Di) * Gt ** 2 / (2 * c["rho"])   # sin codos ni boquillas
        r.update(hi=hi, v_t=Gt / c["rho"], Re_t=Re, Pr_t=Pr, dPt=dPt)
        # ---- Casco: condensación sobre haz horizontal (Nusselt-Kern)
        Nr = max(1.0, 2.0 / 3.0 * Ds / Pt)
        dT, ho = 10.0, 5000.0
        for _ in range(80):
            ho = h_cond_horizontal(h["k_liquido"], h["rho_liquido"], h["rho_vapor"],
                                   h["mu_liquido"], h["lambda"], dT, Do, Nr)
            Uc = 1.0 / (1 / ho + Rw + Do / (Di * hi))
            dT_new = max(Uc * (Tsat - Tc) / ho, 0.5)
            if abs(dT_new - dT) < 1e-4:
                break
            dT = 0.5 * dT + 0.5 * dT_new
        k_ = kern_casco(h["Wh"], h["rho_vapor"], h["mu_vapor"], 1.0, 1.0, Ds, B, Pt, Do, L, arr)
        dPs = 0.5 * Ns * k_["dP"]    # factor 0.5: el caudal de vapor decae al condensar
        r.update(ho=ho, dT_film=dT, dPs=dPs, v_s=k_["v"], Re_s=k_["Re"], Nb=k_["Nb"],
                 De=k_["De"], As=k_["As"], Gs=k_["Gs"], Nr=Nr)
    else:
        # ---- Casco: agua (sensible, Kern)
        k_ = kern_casco(c["Wc"], c["rho"], c["mu"], c["k"], c["cp"], Ds, B, Pt, Do, L, arr)
        ho = k_["hs"]
        dPs = Ns * k_["dP"]
        r.update(ho=ho, dPs=dPs, v_s=k_["v"], Re_s=k_["Re"], Nb=k_["Nb"],
                 De=k_["De"], As=k_["As"], Gs=k_["Gs"])
        # ---- Tubos: condensación en el interior (película laminar vertical)
        dT, hi = 10.0, 5000.0
        for _ in range(80):
            hi = h_cond_vertical(h["k_liquido"], h["rho_liquido"], h["rho_vapor"],
                                 h["mu_liquido"], h["lambda"], dT, L)
            Uc = 1.0 / (1 / ho + Rw + Do / (Di * hi))
            dT_new = max(Uc * (Tc * 0 + Tsat - Tc) * (Do / Di) / hi, 0.5)
            if abs(dT_new - dT) < 1e-4:
                break
            dT = 0.5 * dT + 0.5 * dT_new
        Ap = Nt / n * math.pi * Di ** 2 / 4
        Gt = h["Wh"] / Ap
        Re = Gt * Di / h["mu_vapor"]
        f = factor_friccion(Re, Di)
        dPt = 0.5 * Ns * f * (L * n / Di) * Gt ** 2 / (2 * h["rho_vapor"])
        r.update(hi=hi, dT_film=dT, v_t=Gt / h["rho_vapor"], Re_t=Re, dPt=dPt)

    Ud = d["Q"] / (A * d["DMLT"] * d["Ft"])
    r.update(Uc=Uc, Ud=Ud, Rd_calc=(Uc - Ud) / (Uc * Ud), A_req_final=d["Q"] / (
        (1 / (1 / Uc + d["Rd_req"])) * d["DMLT"] * d["Ft"]))
    dp_t_adm = c["dp_adm"] if d["condensa_en"] == "casco" else h["dp_adm"]
    dp_s_adm = h["dp_adm"] if d["condensa_en"] == "casco" else c["dp_adm"]
    r.update(dp_t_adm=dp_t_adm, dp_s_adm=dp_s_adm,
             ok_dp=(r["dPt"] <= dp_t_adm and r["dPs"] <= dp_s_adm),
             ok_Rd=(r["Rd_calc"] >= d["Rd_req"]))
    return r


def buscar_geometria(d, A_req, Ns, log):
    """Prediseño: selecciona L, pasos, Ds y espaciado de deflectores para un área dada."""
    geo = d["geometria"]
    Do, Pt, arr, rear = geo["Do"], geo["Pt"], geo["arreglo"], d["rear"]
    h, c = d["caliente"], d["fria"]
    isoterma = abs(h["T1"] - h["T2"]) < 1e-9
    n_opts = [1, 2, 4, 6, 8] if (isoterma and rear != "U") else [2, 4, 6, 8]
    L_opts = sorted({l for l in L_STD if l <= geo["L_max"]} | {geo["L_max"]}, reverse=True)
    vmin, vmax = geo["v_min"], geo["v_max"]
    mejor_fallback = None

    for L in L_opts:
        Nt_tot = math.ceil(A_req / (math.pi * Do * L))
        Nt_s = math.ceil(Nt_tot / Ns)
        for n in n_opts:
            Nt = int(math.ceil(Nt_s / n) * n)
            idx0 = next((i for i, Ds in enumerate(DS_STD)
                         if capacidad_tubos(Ds, n, Do, Pt, arr, rear) >= Nt), None)
            if idx0 is None:
                continue
            for idx in range(idx0, min(idx0 + 4, len(DS_STD))):     # permite agrandar el casco
                Ds = DS_STD[idx]
                frac = 0.40
                while frac <= 1.0001:                               # ajuste de deflectores
                    B = max(min(frac * Ds, 1.52), 0.05, Ds / 5)
                    g = dict(Ns=Ns, Nt=Nt, n=n, L=L, Ds=Ds, B=B)
                    r = evaluar(d, g)
                    if r["dPs"] <= r["dp_s_adm"]:
                        break
                    frac += 0.05
                v_ok = True
                if d["condensa_en"] == "casco":
                    v_ok = vmin <= r["v_t"] <= vmax
                if r["ok_dp"] and v_ok:
                    return r
                if r["ok_dp"] and mejor_fallback is None:
                    mejor_fallback = r
    if mejor_fallback:
        log.append("  Aviso: no se encontró geometría con v dentro de [v_min, v_max]; "
                   "se usa la mejor con Δp admisible.")
    return mejor_fallback


def disenar(d):
    log, warns = [], []
    preparar(d, warns)
    h, c = d["caliente"], d["fria"]
    T1, T2, t1, t2 = h["T1"], h["T2"], c["t1"], c["t2"]

    # --- Balance de energía y DMLT
    d["DMLT"] = dmlt_contracorriente(T1, T2, t1, t2)
    log.append(f"Q = {d['Q'] / 1e3:,.1f} kW | Wc = {c['Wc']:.3f} kg/s | DMLT = {d['DMLT']:.2f} °C")

    # --- Factor Ft y cascos en serie
    Ns = 1
    while True:
        Ft = factor_ft(T1, T2, t1, t2, Ns)
        log.append(f"Ns = {Ns} casco(s) en serie -> Ft = {Ft:.3f}")
        if Ft >= 0.75 or Ns >= 6:
            break
        log.append("  Ft < 0.75 -> el agente agrega un casco en serie")
        Ns += 1
    if Ft < 0.75:
        warns.append("Ft < 0.75 incluso con 6 cascos: revisar temperaturas de proceso.")
    d["Ft"], d["Ns"] = Ft, Ns

    # --- Iteración sobre Ud estimado
    Ud_est = d["geometria"]["Ud_est"]
    visitados, factibles, ultimo = set(), [], None
    for it in range(1, 41):
        A_prov = d["Q"] / (Ud_est * d["DMLT"] * Ft)
        r = buscar_geometria(d, A_prov, Ns, log)
        if r is None:
            log.append(f"[{it}] Ud_est={Ud_est:.0f}: sin geometría factible -> se reduce Ud_est")
            Ud_est *= 0.9
            continue
        ultimo = r
        firma = (r["Nt"], r["n"], round(r["L"], 2), round(r["Ds"], 3), round(r["B"], 3))
        exceso = r["A"] / r["A_req_final"] - 1
        log.append(f"[{it}] Ud_est={Ud_est:7.1f} A'={A_prov:6.2f} m2 -> Nt={r['Nt']}x{Ns}, n={r['n']}, "
                   f"L={r['L']:.2f}, Ds={r['Ds'] * 1000:.0f} mm, B={r['B'] * 1000:.0f} mm | "
                   f"hi={r['hi']:.0f} ho={r['ho']:.0f} Uc={r['Uc']:.0f} Ud={r['Ud']:.0f} "
                   f"Rd={r['Rd_calc']:.5f} (req {d['Rd_req']:.5f}) dPt={r['dPt'] / 1e3:.1f} kPa "
                   f"dPs={r['dPs'] / 1e3:.1f} kPa")
        Ud_tgt = 1 / (1 / r["Uc"] + d["Rd_req"])      # Ud máximo que aún cumple Rd_req
        if r["ok_Rd"]:
            factibles.append(r)
            if exceso <= 0.25:
                log.append(f"  Cumple Rd y Δp con exceso de área {exceso * 100:.1f} % -> CONVERGE")
                break
            nuevo = Ud_est * Ud_tgt / r["Ud"] * 0.97
            log.append(f"  Cumple, pero con exceso de área {exceso * 100:.1f} % -> se intenta ajustar")
        else:
            nuevo = Ud_est * Ud_tgt / r["Ud"] * 0.98
            log.append("  Rd_calc < Rd_req -> el agente aumenta el área (reduce Ud_est)")
        if firma in visitados:
            log.append("  Geometría ya evaluada -> se detiene la búsqueda")
            break
        visitados.add(firma)
        Ud_est = nuevo

    final = min(factibles, key=lambda x: x["A"]) if factibles else ultimo
    if final is None:
        raise RuntimeError("No se encontró ninguna geometría factible.")
    if not final["ok_Rd"]:
        warns.append("El diseño final NO cumple Rd_requerido.")
    if not final["ok_dp"]:
        warns.append("El diseño final NO cumple la pérdida de carga admisible.")
    return final, log, warns


# ----------------------------------------------------------------------------
# FASE III: HOJA DE ESPECIFICACIONES TEMA (.xlsx)
# ----------------------------------------------------------------------------
def exportar_hoja_tema(d, r, log, warns, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja TEMA"
    F = "Arial"
    thin = Side(style="thin", color="999999")
    brd = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill_h = PatternFill("solid", fgColor="1F3864")
    fill_s = PatternFill("solid", fgColor="D9E2F3")
    ws.column_dimensions["A"].width = 38
    for col in "BCD":
        ws.column_dimensions[col].width = 24
    ws.column_dimensions["E"].width = 14

    row = [1]
    ref = {}

    def titulo(txt):
        ws.cell(row[0], 1, txt)
        ws.merge_cells(start_row=row[0], start_column=1, end_row=row[0], end_column=5)
        c_ = ws.cell(row[0], 1)
        c_.font = Font(name=F, bold=True, color="FFFFFF", size=11)
        c_.fill = fill_h
        row[0] += 1

    def encabezado(*vals):
        for i, v in enumerate(vals):
            c_ = ws.cell(row[0], 1 + i, v)
            c_.font = Font(name=F, bold=True, size=10)
            c_.fill = fill_s
            c_.border = brd
            c_.alignment = Alignment(horizontal="center")
        row[0] += 1

    def fila(*vals, key=None, fmt=None):
        for i, v in enumerate(vals):
            c_ = ws.cell(row[0], 1 + i, v)
            c_.font = Font(name=F, size=10)
            c_.border = brd
            if i > 0 and fmt and isinstance(v, (int, float, str)) and (not isinstance(v, str) or v.startswith("=")):
                c_.number_format = fmt
            if i > 0:
                c_.alignment = Alignment(horizontal="right")
        if key:
            ref[key] = row[0]
        row[0] += 1

    idn, h, c, geo = d["identificacion"], d["caliente"], d["fria"], d["geometria"]
    cas = d["condensa_en"] == "casco"
    S = (h if cas else c)   # fluido lado casco
    T = (c if cas else h)   # fluido lado tubos

    ws.cell(row[0], 1, "HOJA DE ESPECIFICACIONES TÉCNICAS — INTERCAMBIADOR DE CASCO Y TUBOS (TEMA)")
    ws.merge_cells(start_row=row[0], start_column=1, end_row=row[0], end_column=5)
    ws.cell(row[0], 1).font = Font(name=F, bold=True, size=13)
    row[0] += 2

    # 1. Identificación
    titulo("1. DATOS DE IDENTIFICACIÓN DEL EQUIPO")
    fila("Ítem / Tag", idn["item"])
    fila("Servicio", idn["servicio"])
    fila("Cliente / Proyecto", idn["cliente"])
    fila("Tipo TEMA", idn["tema"])
    fila("Orientación", idn["orientacion"])
    fila("Cascos en serie", r["Ns"])
    row[0] += 1

    # 2. Condiciones de operación
    titulo("2. CONDICIONES DE OPERACIÓN Y PROCESO")
    encabezado("Parámetro", "Lado casco", "Lado tubos", "Unidad")
    fila("Fluido", S["fluido"], T["fluido"], "-")
    wS = h["Wh"] if cas else c["Wc"]
    wT = c["Wc"] if cas else h["Wh"]
    fila("Caudal másico", round(wS, 3), round(wT, 3), "kg/s", fmt="0.000")
    fila("Temperatura de entrada", S.get("T1", S.get("t1")), T.get("T1", T.get("t1")), "°C", fmt="0.0")
    fila("Temperatura de salida", S.get("T2", S.get("t2")), T.get("T2", T.get("t2")), "°C", fmt="0.0")
    fila("Presión de operación", round(S["P_oper"] / 1e3, 1), round(T["P_oper"] / 1e3, 1), "kPa abs", fmt="0.0")
    if cas:
        fila("Calor latente de condensación", round(h["lambda"] / 1e3, 1), "-", "kJ/kg", fmt="0.0")
        fila("Densidad (líquido / vapor)", f"{h['rho_liquido']} / {h['rho_vapor']}", c["rho"], "kg/m³")
        fila("Viscosidad", f"{h['mu_liquido'] * 1e3:.4f}", f"{c['mu'] * 1e3:.4f}", "mPa·s")
        fila("Conductividad térmica", h["k_liquido"], c["k"], "W/m·K")
        fila("Calor específico", "-", c["cp"], "J/kg·K")
    else:
        fila("Calor latente de condensación", "-", round(h["lambda"] / 1e3, 1), "kJ/kg", fmt="0.0")
        fila("Densidad", c["rho"], f"{h['rho_liquido']} / {h['rho_vapor']}", "kg/m³")
        fila("Viscosidad", f"{c['mu'] * 1e3:.4f}", f"{h['mu_liquido'] * 1e3:.4f}", "mPa·s")
        fila("Conductividad térmica", c["k"], h["k_liquido"], "W/m·K")
        fila("Calor específico", c["cp"], "-", "J/kg·K")
    fila("Velocidad", round(r["v_s"], 2), round(r["v_t"], 2), "m/s", fmt="0.00")
    fila("Factor de ensuciamiento", S["Rd"], T["Rd"], "m²·K/W", fmt="0.00000")
    fila("Δp admisible", round(r["dp_s_adm"] / 1e3, 1), round(r["dp_t_adm"] / 1e3, 1), "kPa", fmt="0.0")
    fila("Δp calculada (sin codos ni boquillas)", round(r["dPs"] / 1e3, 2), round(r["dPt"] / 1e3, 2), "kPa", fmt="0.00")
    fila("Carga térmica Q", round(d["Q"] / 1e3, 1), "", "kW", fmt="#,##0.0")
    fila("DMLT (contracorriente)", round(d["DMLT"], 2), "", "°C", fmt="0.00")
    fila("Factor Ft", round(d["Ft"], 3), "", "-", fmt="0.000")
    row[0] += 1

    # 3. Geometría
    titulo("3. DATOS GEOMÉTRICOS Y CONSTRUCTIVOS")
    encabezado("Parámetro", "Haz tubular", "Carcasa", "Unidad")
    fila("Diámetro externo / interno de tubo", f"{geo['Do'] * 1000:.2f} / {geo['Di'] * 1000:.2f}", "", "mm")
    fila("Diámetro interno de carcasa Ds", "", round(r["Ds"] * 1000, 1), "mm", fmt="0.0")
    fila("Número de tubos por casco (Nt)", r["Nt"], "", "-", key="Nt")
    fila("Longitud de tubo (L)", r["L"], "", "m", key="L", fmt="0.00")
    fila("Número de pasos por tubo", r["n"], "", "-")
    fila("Arreglo de tubos", "Triangular 60°" if geo["arreglo"] != "cuadro" else "Cuadro 90°", "", "-")
    fila("Paso de tubos (Pt)", round(geo["Pt"] * 1000, 2), "", "mm", fmt="0.00")
    fila("Material", idn["material_tubos"], idn["material_casco"], "-")
    fila("Espaciado de deflectores (B)", "", round(r["B"] * 1000, 1), "mm", fmt="0.0")
    fila("Número de deflectores", "", r["Nb"], "-")
    fila("Corte de deflector", "", "25 % segmentado", "-")
    fila("Superficie por casco", "", "", "m²")
    ref["A_casco"] = row[0] - 1
    row[0] += 1

    # 4. Resultados
    titulo("4. RESULTADOS DE DESEMPEÑO TÉRMICO E HIDRÁULICO")
    encabezado("Parámetro", "Valor", "", "Unidad")
    fila("ho (coeficiente película lado casco)", round(r["ho"], 1), "", "W/m²·K", fmt="#,##0.0")
    fila("hi (coeficiente película lado tubos)", round(r["hi"], 1), "", "W/m²·K", fmt="#,##0.0")
    fila("Uc (limpio)", round(r["Uc"], 1), "", "W/m²·K", key="Uc", fmt="#,##0.0")
    fila("Ud (diseño)", round(r["Ud"], 1), "", "W/m²·K", key="Ud", fmt="#,##0.0")
    fila("Rd_calc = (Uc - Ud)/(Uc·Ud)", "", "", "m²·K/W", key="Rd")
    fila("Rd_requerido", round(d["Rd_req"], 6), "", "m²·K/W", key="Rdreq", fmt="0.000000")
    fila("Área de intercambio requerida", round(r["A_req_final"], 2), "", "m²", key="Areq", fmt="0.00")
    fila("Área de intercambio proporcionada", "", "", "m²", key="Aprov")
    fila("Exceso de área", "", "", "%", key="Exc")
    fila("Δp calculada lado casco", round(r["dPs"] / 1e3, 2), "", "kPa", fmt="0.00")
    fila("Δp calculada lado tubos", round(r["dPt"] / 1e3, 2), "", "kPa", fmt="0.00")
    fila("Cumple Rd y Δp", "SÍ" if (r["ok_Rd"] and r["ok_dp"]) else "NO", "", "-")

    # Fórmulas vivas (se recalculan si se editan Nt, L, Uc, Ud...)
    ws.cell(ref["A_casco"], 3, f"=B{ref['Nt']}*PI()*{geo['Do']}*B{ref['L']}")
    ws.cell(ref["A_casco"], 3).number_format = "0.00"
    ws.cell(ref["Rd"], 2, f"=(B{ref['Uc']}-B{ref['Ud']})/(B{ref['Uc']}*B{ref['Ud']})")
    ws.cell(ref["Rd"], 2).number_format = "0.000000"
    ws.cell(ref["Aprov"], 2, f"={r['Ns']}*C{ref['A_casco']}")
    ws.cell(ref["Aprov"], 2).number_format = "0.00"
    ws.cell(ref["Exc"], 2, f"=(B{ref['Aprov']}/B{ref['Areq']}-1)*100")
    ws.cell(ref["Exc"], 2).number_format = "0.0"
    for k_ in ("A_casco", "Rd", "Aprov", "Exc"):
        for col in (2, 3):
            ws.cell(ref[k_], col).font = Font(name=F, size=10)
            ws.cell(ref[k_], col).alignment = Alignment(horizontal="right")

    row[0] += 1
    ws.cell(row[0], 1, "Notas: Q = Wh·λ; coef. de condensación por Nusselt/Kern; lado sensible Gnielinski/Kern; "
                       "Δp sin codos ni boquillas. Δp con vapor condensante afectada por factor 0.5.")
    ws.cell(row[0], 1).font = Font(name=F, italic=True, size=9)
    for w_ in warns:
        row[0] += 1
        ws.cell(row[0], 1, "AVISO: " + w_).font = Font(name=F, bold=True, color="C00000", size=9)

    # Hoja con historial de iteraciones
    ws2 = wb.create_sheet("Iteraciones")
    ws2.column_dimensions["A"].width = 160
    ws2.cell(1, 1, "Registro de decisiones del agente").font = Font(name=F, bold=True, size=11)
    for i, linea in enumerate(log, start=2):
        ws2.cell(i, 1, linea).font = Font(name=F, size=9)

    wb.save(ruta)


# ----------------------------------------------------------------------------
# PROGRAMA PRINCIPAL
# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--entrada", help="JSON con los datos de proceso (Fase I)")
    ap.add_argument("--salida", default="hoja_TEMA.xlsx")
    ap.add_argument("--ejemplo", help="Guarda un JSON de ejemplo en esa ruta y termina")
    ap.add_argument("--editar", action="store_true", help="Abre los datos de entrada en $EDITOR para poder modificarlos antes de ejecutar")
    a = ap.parse_args()

    if a.ejemplo:
        with open(a.ejemplo, "w", encoding="utf-8") as f:
            json.dump(EJEMPLO, f, indent=2, ensure_ascii=False)
        print("Ejemplo guardado en", a.ejemplo)
        return

    if a.entrada:
        with open(a.entrada, encoding="utf-8") as f:
            d = json.load(f)
    else:
        d = json.loads(json.dumps(EJEMPLO))

    # Si se pide editar, volcar los datos a un archivo temporal y abrir el editor del usuario
    if a.editar:
        import tempfile, os, subprocess
        fd, tmp = tempfile.mkstemp(suffix='.json')
        os.close(fd)
        with open(tmp, 'w', encoding='utf-8') as f:
            json.dump(d, f, indent=2, ensure_ascii=False)
        editor = os.environ.get('EDITOR', 'nano')
        try:
            subprocess.call([editor, tmp])
            with open(tmp, encoding='utf-8') as f:
                d = json.load(f)
        finally:
            try:
                os.remove(tmp)
            except Exception:
                pass

    r, log, warns = disenar(d)
    print("\n".join(log))
    print("\n===== RESULTADO FINAL =====")
    print(f"Casco(s): {r['Ns']} en serie | Ds = {r['Ds'] * 1000:.0f} mm | B = {r['B'] * 1000:.0f} mm")
    print(f"Tubos por casco: {r['Nt']} | L = {r['L']:.2f} m | pasos = {r['n']}")
    print(f"ho = {r['ho']:.0f} | hi = {r['hi']:.0f} W/m2K | Uc = {r['Uc']:.0f} | Ud = {r['Ud']:.0f}")
    print(f"Rd_calc = {r['Rd_calc']:.6f} (req {d['Rd_req']:.6f}) | Área = {r['A']:.2f} m2")
    print(f"dP casco = {r['dPs'] / 1e3:.2f} kPa | dP tubos = {r['dPt'] / 1e3:.2f} kPa")
    for w in warns:
        print("AVISO:", w)
    exportar_hoja_tema(d, r, log, warns, a.salida)
    print("Hoja TEMA exportada en", a.salida)


if __name__ == "__main__":
    main()
